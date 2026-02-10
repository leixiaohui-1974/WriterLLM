"""End-to-end pipeline integration tests."""
import os
import shutil
import io
import pytest
import docx

from src.parser import parse_document
from src.generator import generate_slides
from src.models import SlideData, SlideLayout, SlideTheme, Language, ExportFormat, PresentationConfig
from src.renderer import (
    create_pptx_file, create_slide_images, create_pdf_from_images,
    _prepare_background, _get_font, _is_cjk_char, _contains_cjk,
    _wrap_text, _text_pixel_width, _draw_text,
    _compute_content_font_size, _estimate_content_lines, _cover_crop,
    CONTENT_MAX_WIDTH, CONTENT_FONT_SIZE, FOOTER_AREA, CONTENT_START_Y,
    SLIDE_HEIGHT,
)
from src.image_gen import generate_slide_image, generate_slide_images_batch
from src.models import ThemeColors, THEMES
from src.generator import mock_generate_content


OUTPUT_DIR = "test_output"


@pytest.fixture(autouse=True)
def clean_output():
    """Clean test output directory before and after each test."""
    if os.path.exists(OUTPUT_DIR):
        shutil.rmtree(OUTPUT_DIR)
    os.makedirs(OUTPUT_DIR)
    yield
    if os.path.exists(OUTPUT_DIR):
        shutil.rmtree(OUTPUT_DIR)


def _create_test_docx():
    """Create a test DOCX document in memory."""
    doc = docx.Document()
    doc.add_heading("AI in Healthcare", 0)
    doc.add_paragraph("AI is revolutionizing healthcare diagnostics.")
    doc.add_paragraph("Machine learning models can predict diseases early.")
    doc.add_paragraph("Robotic surgery is becoming more precise.")
    doc.add_paragraph("Drug discovery is accelerated by AI algorithms.")
    doc.add_paragraph("Patient monitoring systems use AI for early detection.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    buf.name = "test_doc.docx"
    return buf


def _create_test_bg_image(path):
    """Create a test background image."""
    from PIL import Image
    img = Image.new("RGB", (1920, 1080), color=(100, 150, 200))
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    img.save(path, "PNG")
    return path


class TestFullPipeline:
    def test_parse_generate_render(self):
        """Test the complete pipeline: parse -> generate -> render."""
        doc_io = _create_test_docx()
        text = parse_document(doc_io)
        assert "AI is revolutionizing" in text

        slides_data = generate_slides(text, num_slides=3)
        assert len(slides_data) == 3
        for s in slides_data:
            assert isinstance(s, SlideData)

        pptx_path = os.path.join(OUTPUT_DIR, "test.pptx")
        create_pptx_file(slides_data, pptx_path)
        assert os.path.exists(pptx_path)
        assert os.path.getsize(pptx_path) > 0

        images_dir = os.path.join(OUTPUT_DIR, "images")
        image_paths = create_slide_images(slides_data, images_dir)
        assert len(image_paths) == 3
        for p in image_paths:
            assert os.path.exists(p)
            assert os.path.getsize(p) > 0

        pdf_path = os.path.join(OUTPUT_DIR, "test.pdf")
        create_pdf_from_images(image_paths, pdf_path)
        assert os.path.exists(pdf_path)
        assert os.path.getsize(pdf_path) > 0

    def test_themed_pptx_generation(self):
        """Test that themed PPTX files are created successfully."""
        slides_data = generate_slides("Test content. " * 20, num_slides=3)
        for theme in SlideTheme:
            path = os.path.join(OUTPUT_DIR, f"test_{theme.value}.pptx")
            create_pptx_file(slides_data, path, theme=theme)
            assert os.path.exists(path)
            assert os.path.getsize(path) > 0


class TestThemeRendering:
    def test_all_themes_render(self):
        """Verify that all themes produce valid output."""
        slides_data = generate_slides("Test content for theme rendering. " * 20, num_slides=2)
        for theme in SlideTheme:
            images_dir = os.path.join(OUTPUT_DIR, f"images_{theme.value}")
            image_paths = create_slide_images(slides_data, images_dir, theme=theme)
            assert len(image_paths) == 2
            for p in image_paths:
                assert os.path.exists(p)
                assert os.path.getsize(p) > 0


class TestLayoutRendering:
    def test_all_layout_types_render(self):
        """Verify that all layout types render correctly."""
        slides_data = [
            SlideData(title="Title Slide", content=["Subtitle here"], layout=SlideLayout.TITLE,
                      speaker_notes="Welcome", image_prompt="Title visual"),
            SlideData(title="Content Slide", content=["Point 1", "Point 2", "Point 3"], layout=SlideLayout.CONTENT,
                      speaker_notes="Content notes", image_prompt="Content visual"),
            SlideData(title="Section Break", content=[], layout=SlideLayout.SECTION,
                      speaker_notes="Section notes", image_prompt="Section visual"),
            SlideData(title="Comparison", content=["Left 1", "Left 2", "Right 1", "Right 2"],
                      layout=SlideLayout.TWO_COLUMN, speaker_notes="Comparison notes", image_prompt="Compare visual"),
        ]
        images_dir = os.path.join(OUTPUT_DIR, "images_layouts")
        image_paths = create_slide_images(slides_data, images_dir)
        assert len(image_paths) == 4
        for p in image_paths:
            assert os.path.exists(p)
            assert os.path.getsize(p) > 0

    def test_layout_pptx_rendering(self):
        """Verify PPTX rendering with different layouts."""
        slides_data = [
            SlideData(title="Title Slide", content=["Subtitle"], layout=SlideLayout.TITLE,
                      speaker_notes="Welcome"),
            SlideData(title="Content", content=["A", "B"], layout=SlideLayout.CONTENT,
                      speaker_notes="Content notes"),
            SlideData(title="Section", content=[], layout=SlideLayout.SECTION,
                      speaker_notes="Section notes"),
        ]
        pptx_path = os.path.join(OUTPUT_DIR, "layouts.pptx")
        create_pptx_file(slides_data, pptx_path)
        assert os.path.exists(pptx_path)
        assert os.path.getsize(pptx_path) > 0


class TestBackgroundImageCompositing:
    """Tests for AI background image compositing into slides."""

    def test_slide_images_with_background(self):
        """Verify that background images are composited into slide images."""
        slides_data = [
            SlideData(title="Slide 1", content=["Content A"], layout=SlideLayout.CONTENT),
            SlideData(title="Slide 2", content=["Content B"], layout=SlideLayout.TITLE),
        ]
        # Create a fake background image
        bg_path = os.path.join(OUTPUT_DIR, "bg_000.png")
        _create_test_bg_image(bg_path)

        background_images = {0: bg_path}  # Only slide 0 has a background

        images_dir = os.path.join(OUTPUT_DIR, "images_bg")
        image_paths = create_slide_images(
            slides_data, images_dir, background_images=background_images,
        )
        assert len(image_paths) == 2
        for p in image_paths:
            assert os.path.exists(p)
            # Files with bg should have content
            assert os.path.getsize(p) > 0

    def test_pptx_with_background(self):
        """Verify PPTX generation with background images."""
        slides_data = [
            SlideData(title="BG Slide", content=["Test"], layout=SlideLayout.CONTENT),
        ]
        bg_path = os.path.join(OUTPUT_DIR, "bg_pptx.png")
        _create_test_bg_image(bg_path)

        pptx_path = os.path.join(OUTPUT_DIR, "bg_test.pptx")
        create_pptx_file(
            slides_data, pptx_path,
            background_images={0: bg_path},
        )
        assert os.path.exists(pptx_path)
        assert os.path.getsize(pptx_path) > 0

    def test_prepare_background_with_image(self):
        """Test _prepare_background composites correctly."""
        bg_path = os.path.join(OUTPUT_DIR, "bg_prep.png")
        _create_test_bg_image(bg_path)

        colors = THEMES[SlideTheme.PROFESSIONAL]
        img = _prepare_background(0, colors, background_images={0: bg_path})
        assert img.size == (1920, 1080)
        assert img.mode == "RGB"

    def test_prepare_background_without_image(self):
        """Test _prepare_background returns solid color without image."""
        colors = THEMES[SlideTheme.DARK]
        img = _prepare_background(0, colors, background_images=None)
        assert img.size == (1920, 1080)
        # Should be the dark theme background color
        pixel = img.getpixel((100, 100))
        assert pixel == colors.background

    def test_all_layouts_with_background(self):
        """Verify all layout types render correctly with background images."""
        slides_data = [
            SlideData(title="Title", content=["Sub"], layout=SlideLayout.TITLE),
            SlideData(title="Content", content=["A", "B"], layout=SlideLayout.CONTENT),
            SlideData(title="Section", content=[], layout=SlideLayout.SECTION),
            SlideData(title="Two Col", content=["L1", "L2", "R1", "R2"], layout=SlideLayout.TWO_COLUMN),
        ]
        bg_dir = os.path.join(OUTPUT_DIR, "bgs")
        os.makedirs(bg_dir, exist_ok=True)
        bg_images = {}
        for i in range(4):
            bg_path = os.path.join(bg_dir, f"bg_{i}.png")
            _create_test_bg_image(bg_path)
            bg_images[i] = bg_path

        images_dir = os.path.join(OUTPUT_DIR, "images_all_bg")
        image_paths = create_slide_images(
            slides_data, images_dir, background_images=bg_images,
        )
        assert len(image_paths) == 4
        for p in image_paths:
            assert os.path.exists(p)


class TestCJKTextWrapping:
    """Tests for CJK character detection and pixel-based text wrapping."""

    def test_is_cjk_char_chinese(self):
        assert _is_cjk_char("人") is True
        assert _is_cjk_char("智") is True

    def test_is_cjk_char_japanese(self):
        assert _is_cjk_char("あ") is True  # Hiragana
        assert _is_cjk_char("カ") is True  # Katakana

    def test_is_cjk_char_korean(self):
        assert _is_cjk_char("한") is True  # Hangul

    def test_is_cjk_char_latin(self):
        assert _is_cjk_char("A") is False
        assert _is_cjk_char("z") is False
        assert _is_cjk_char(" ") is False

    def test_contains_cjk(self):
        assert _contains_cjk("人工智能") is True
        assert _contains_cjk("Hello World") is False
        assert _contains_cjk("Hello 世界") is True  # Mixed
        assert _contains_cjk("") is False

    def test_wrap_text_latin(self):
        font = _get_font("DejaVuSans.ttf", 45, Language.ENGLISH)
        result = _wrap_text("Short text", font, 1720)
        assert len(result) >= 1
        assert "Short text" in result[0]

    def test_wrap_text_cjk_wraps_long_string(self):
        """CJK text without spaces must be wrapped by pixel width."""
        font = _get_font("DejaVuSans.ttf", 45, Language.CHINESE)
        long_cjk = "人工智能正在改变世界的方方面面包括医疗金融教育交通等各个领域"
        result = _wrap_text(long_cjk, font, 800)  # Narrow width
        assert len(result) >= 2, f"Expected multiple lines, got {len(result)}: {result}"

    def test_wrap_text_cjk_short_fits(self):
        """Short CJK text should fit in one line."""
        font = _get_font("DejaVuSans.ttf", 45, Language.CHINESE)
        result = _wrap_text("人工智能", font, 1720)
        assert len(result) == 1

    def test_wrap_text_empty(self):
        font = _get_font("DejaVuSans.ttf", 45, Language.ENGLISH)
        assert _wrap_text("", font, 1000) == []
        assert _wrap_text("   ", font, 1000) == []

    def test_wrap_text_mixed_cjk_latin(self):
        """Mixed CJK+Latin text should wrap correctly."""
        font = _get_font("DejaVuSans.ttf", 45, Language.CHINESE)
        mixed = "AI人工智能is changing the world世界"
        result = _wrap_text(mixed, font, 600)
        assert len(result) >= 1

    def test_text_pixel_width(self):
        font = _get_font("DejaVuSans.ttf", 45, Language.ENGLISH)
        assert _text_pixel_width("", font) == 0
        w = _text_pixel_width("Hello", font)
        assert w > 0
        # Longer text should be wider
        w2 = _text_pixel_width("Hello World", font)
        assert w2 > w


class TestCJKFontSupport:
    """Tests for CJK font loading and rendering."""

    def test_cjk_font_loading_chinese(self):
        font = _get_font("DejaVuSans.ttf", 45, Language.CHINESE)
        assert font is not None

    def test_cjk_font_loading_japanese(self):
        font = _get_font("DejaVuSans.ttf", 45, Language.JAPANESE)
        assert font is not None

    def test_cjk_font_loading_korean(self):
        font = _get_font("DejaVuSans.ttf", 45, Language.KOREAN)
        assert font is not None

    def test_english_font_loading(self):
        font = _get_font("DejaVuSans.ttf", 45, Language.ENGLISH)
        assert font is not None

    def test_chinese_slide_rendering(self):
        """Verify Chinese text renders without error."""
        slides_data = [
            SlideData(title="人工智能", content=["机器学习", "深度学习"], layout=SlideLayout.CONTENT),
        ]
        images_dir = os.path.join(OUTPUT_DIR, "images_zh")
        image_paths = create_slide_images(
            slides_data, images_dir, language=Language.CHINESE,
        )
        assert len(image_paths) == 1
        assert os.path.getsize(image_paths[0]) > 0

    def test_japanese_slide_rendering(self):
        """Verify Japanese text renders without error."""
        slides_data = [
            SlideData(title="プレゼンテーション", content=["ポイント1", "ポイント2"], layout=SlideLayout.CONTENT),
        ]
        images_dir = os.path.join(OUTPUT_DIR, "images_ja")
        image_paths = create_slide_images(
            slides_data, images_dir, language=Language.JAPANESE,
        )
        assert len(image_paths) == 1
        assert os.path.getsize(image_paths[0]) > 0

    def test_chinese_section_layout(self):
        """Verify section layout renders CJK title with correct font."""
        slides_data = [
            SlideData(title="第二部分", content=[], layout=SlideLayout.SECTION),
        ]
        images_dir = os.path.join(OUTPUT_DIR, "images_zh_section")
        image_paths = create_slide_images(
            slides_data, images_dir, language=Language.CHINESE,
        )
        assert len(image_paths) == 1
        assert os.path.getsize(image_paths[0]) > 0

    def test_cjk_long_content_wrapping(self):
        """Verify long Chinese text wraps correctly in slide images."""
        long_text = "人工智能正在改变世界的方方面面包括医疗金融教育交通等各个领域，它能够帮助我们更好地理解和分析大量数据"
        slides_data = [
            SlideData(title="长文本测试", content=[long_text], layout=SlideLayout.CONTENT),
        ]
        images_dir = os.path.join(OUTPUT_DIR, "images_zh_long")
        image_paths = create_slide_images(
            slides_data, images_dir, language=Language.CHINESE,
        )
        assert len(image_paths) == 1
        assert os.path.getsize(image_paths[0]) > 0


class TestTextShadow:
    """Tests for text shadow rendering with background images."""

    def test_slide_with_bg_renders_shadow(self):
        """Verify slides with background images render without error (shadow active)."""
        slides_data = [
            SlideData(title="Shadow Test", content=["Point A", "Point B"], layout=SlideLayout.CONTENT),
        ]
        bg_path = os.path.join(OUTPUT_DIR, "bg_shadow.png")
        _create_test_bg_image(bg_path)

        images_dir = os.path.join(OUTPUT_DIR, "images_shadow")
        image_paths = create_slide_images(
            slides_data, images_dir, background_images={0: bg_path},
        )
        assert len(image_paths) == 1
        assert os.path.getsize(image_paths[0]) > 0

    def test_draw_text_with_shadow(self):
        """Verify _draw_text with shadow draws more dark pixels than without."""
        from PIL import Image, ImageDraw
        # Render without shadow
        img_no_shadow = Image.new("RGB", (400, 100), (255, 255, 255))
        draw1 = ImageDraw.Draw(img_no_shadow)
        font = _get_font("DejaVuSans.ttf", 30, Language.ENGLISH)
        _draw_text(draw1, (10, 10), "Shadow", font, (255, 0, 0), shadow=False)
        px1 = img_no_shadow.load()
        dark_no_shadow = sum(1 for x in range(400) for y in range(100) if px1[x, y][0] < 200)

        # Render with shadow
        img_shadow = Image.new("RGB", (400, 100), (255, 255, 255))
        draw2 = ImageDraw.Draw(img_shadow)
        _draw_text(draw2, (10, 10), "Shadow", font, (255, 0, 0), shadow=True)
        px2 = img_shadow.load()
        dark_shadow = sum(1 for x in range(400) for y in range(100) if px2[x, y][0] < 200)

        # Shadow version should have more dark pixels (the black shadow layer)
        assert dark_shadow > dark_no_shadow, "Shadow should add more dark pixels"


class TestMultiLanguage:
    def test_chinese_content_generation(self):
        text = "人工智能正在改变世界。它在医疗和金融领域广泛应用。机器学习是人工智能的重要分支。"
        slides = generate_slides(text, num_slides=2, language=Language.CHINESE)
        assert len(slides) == 2

    def test_english_content_generation(self):
        text = "Artificial Intelligence is transforming the world. It is used in many fields."
        slides = generate_slides(text, num_slides=2, language=Language.ENGLISH)
        assert len(slides) == 2


class TestModels:
    def test_slide_data_to_dict(self):
        slide = SlideData(title="Test", content=["A", "B"], speaker_notes="Notes", image_prompt="Prompt")
        d = slide.to_dict()
        assert d["title"] == "Test"
        assert d["content"] == ["A", "B"]
        assert "layout" in d

    def test_slide_data_from_dict(self):
        d = {"title": "Test", "content": ["A"], "speaker_notes": "Notes", "image_prompt": "Prompt"}
        slide = SlideData.from_dict(d)
        assert slide.title == "Test"
        assert slide.content == ["A"]

    def test_slide_data_from_dict_defaults(self):
        slide = SlideData.from_dict({})
        assert slide.title == "Untitled"
        assert slide.content == []
        assert slide.layout == SlideLayout.CONTENT

    def test_slide_data_from_dict_with_layout(self):
        d = {"title": "Test", "layout": "title"}
        slide = SlideData.from_dict(d)
        assert slide.layout == SlideLayout.TITLE

    def test_slide_data_from_dict_invalid_layout(self):
        d = {"title": "Test", "layout": "nonexistent"}
        slide = SlideData.from_dict(d)
        assert slide.layout == SlideLayout.CONTENT

    def test_presentation_config_defaults(self):
        config = PresentationConfig()
        assert config.num_slides == 5
        assert config.language == Language.ENGLISH
        assert ExportFormat.PPTX in config.export_formats
        assert ExportFormat.VIDEO in config.export_formats

    def test_presentation_config_voice_name(self):
        config = PresentationConfig(language=Language.CHINESE, voice_gender="Female")
        assert "zh-CN" in config.voice_name

    def test_export_format_enum(self):
        assert ExportFormat.PPTX.value == "pptx"
        assert ExportFormat.PDF.value == "pdf"
        assert ExportFormat.VIDEO.value == "video"


class TestImageGen:
    def test_no_api_key_returns_none(self):
        result = generate_slide_image("test prompt", "/tmp/test.png")
        assert result is None

    def test_batch_no_api_key_returns_empty(self):
        prompts = [(0, "prompt 1"), (1, "prompt 2")]
        results = generate_slide_images_batch(prompts, OUTPUT_DIR)
        assert results == {}


class TestPdfGeneration:
    def test_empty_image_list_raises(self):
        with pytest.raises(ValueError, match="No images"):
            create_pdf_from_images([], os.path.join(OUTPUT_DIR, "empty.pdf"))


class TestResponsiveFontSizing:
    """Tests for responsive font sizing when content overflows."""

    def test_compute_font_size_short_content(self):
        """Short content should use the base font size."""
        content = ["Point 1", "Point 2"]
        available_h = SLIDE_HEIGHT - CONTENT_START_Y - FOOTER_AREA
        size = _compute_content_font_size(content, Language.ENGLISH, CONTENT_MAX_WIDTH, available_h)
        assert size == CONTENT_FONT_SIZE

    def test_compute_font_size_long_content_shrinks(self):
        """Many bullet points should trigger font shrinking."""
        content = [f"This is bullet point number {i} with enough text to be meaningful" for i in range(15)]
        available_h = SLIDE_HEIGHT - CONTENT_START_Y - FOOTER_AREA
        size = _compute_content_font_size(content, Language.ENGLISH, CONTENT_MAX_WIDTH, available_h)
        assert size < CONTENT_FONT_SIZE
        assert size >= 24  # Should not go below minimum

    def test_compute_font_size_minimum_bound(self):
        """Even very long content should not go below minimum font size."""
        content = [f"Very long line {i}" for i in range(50)]
        available_h = 200  # Very limited height
        size = _compute_content_font_size(content, Language.ENGLISH, CONTENT_MAX_WIDTH, available_h)
        assert size >= 24

    def test_estimate_content_lines(self):
        """Verify line estimation works for simple content."""
        font = _get_font("DejaVuSans.ttf", 45, Language.ENGLISH)
        lines = _estimate_content_lines(["Short", "Also short"], font, CONTENT_MAX_WIDTH)
        assert lines == 2

    def test_estimate_content_lines_wrapping(self):
        """Long text should produce more estimated lines."""
        font = _get_font("DejaVuSans.ttf", 45, Language.ENGLISH)
        long_text = "This is a very long sentence that should definitely wrap around multiple times " * 3
        lines = _estimate_content_lines([long_text], font, CONTENT_MAX_WIDTH)
        assert lines >= 2

    def test_responsive_rendering_many_bullets(self):
        """Verify slide with many bullets renders without error using smaller font."""
        slides_data = [
            SlideData(
                title="Dense Content",
                content=[f"Bullet point {i}: description of item" for i in range(12)],
                layout=SlideLayout.CONTENT,
            ),
        ]
        images_dir = os.path.join(OUTPUT_DIR, "images_responsive")
        image_paths = create_slide_images(slides_data, images_dir)
        assert len(image_paths) == 1
        assert os.path.getsize(image_paths[0]) > 0

    def test_responsive_two_column_rendering(self):
        """Verify two-column layout with many items uses responsive sizing."""
        slides_data = [
            SlideData(
                title="Dense Two Column",
                content=[f"Item {i}" for i in range(14)],
                layout=SlideLayout.TWO_COLUMN,
            ),
        ]
        images_dir = os.path.join(OUTPUT_DIR, "images_responsive_2col")
        image_paths = create_slide_images(slides_data, images_dir)
        assert len(image_paths) == 1
        assert os.path.getsize(image_paths[0]) > 0


class TestCoverCrop:
    """Tests for aspect-ratio-preserving background image resizing."""

    def test_cover_crop_landscape(self):
        from PIL import Image
        img = Image.new("RGBA", (2000, 1000), (100, 100, 100, 255))
        cropped = _cover_crop(img, 1920, 1080)
        assert cropped.size == (1920, 1080)

    def test_cover_crop_portrait(self):
        from PIL import Image
        img = Image.new("RGBA", (800, 1200), (100, 100, 100, 255))
        cropped = _cover_crop(img, 1920, 1080)
        assert cropped.size == (1920, 1080)

    def test_cover_crop_square(self):
        from PIL import Image
        img = Image.new("RGBA", (1000, 1000), (100, 100, 100, 255))
        cropped = _cover_crop(img, 1920, 1080)
        assert cropped.size == (1920, 1080)


class TestConfigurableOverlay:
    """Tests for configurable background overlay opacity."""

    def test_overlay_opacity_default(self):
        """Default overlay opacity should work."""
        bg_path = os.path.join(OUTPUT_DIR, "bg_opacity.png")
        _create_test_bg_image(bg_path)
        colors = THEMES[SlideTheme.PROFESSIONAL]
        img = _prepare_background(0, colors, {0: bg_path})
        assert img.size == (1920, 1080)

    def test_overlay_opacity_transparent(self):
        """Fully transparent overlay should show more of the original image."""
        bg_path = os.path.join(OUTPUT_DIR, "bg_transparent.png")
        _create_test_bg_image(bg_path)
        colors = THEMES[SlideTheme.PROFESSIONAL]
        img_transparent = _prepare_background(0, colors, {0: bg_path}, overlay_opacity=0)
        img_opaque = _prepare_background(0, colors, {0: bg_path}, overlay_opacity=200)
        # Both should be valid images
        assert img_transparent.size == (1920, 1080)
        assert img_opaque.size == (1920, 1080)
        # Transparent should be closer to original BG color (100,150,200)
        px_t = img_transparent.getpixel((500, 500))
        px_o = img_opaque.getpixel((500, 500))
        # More opaque = closer to theme background (255,255,255 for Professional)
        assert px_o[0] > px_t[0] or px_o[1] > px_t[1]

    def test_overlay_opacity_in_slide_images(self):
        """Verify overlay_opacity parameter works in create_slide_images."""
        slides_data = [
            SlideData(title="Opacity Test", content=["Test"], layout=SlideLayout.CONTENT),
        ]
        bg_path = os.path.join(OUTPUT_DIR, "bg_slide_opacity.png")
        _create_test_bg_image(bg_path)
        images_dir = os.path.join(OUTPUT_DIR, "images_opacity")
        image_paths = create_slide_images(
            slides_data, images_dir,
            background_images={0: bg_path},
            overlay_opacity=100,
        )
        assert len(image_paths) == 1
        assert os.path.getsize(image_paths[0]) > 0


class TestMockSummarySlide:
    """Tests for improved mock mode with summary slide generation."""

    def test_mock_generates_summary_slide(self):
        """Mock mode with 5+ slides should produce a summary slide."""
        text = "AI is transforming healthcare. Machine learning improves diagnostics. " * 20
        slides = mock_generate_content(text, num_slides=5)
        assert len(slides) == 5
        last = slides[-1]
        assert last.title == "Key Takeaways"

    def test_mock_summary_chinese(self):
        """Chinese mock mode should produce Chinese summary title."""
        text = "人工智能正在改变世界。机器学习改善诊断。" * 10
        slides = mock_generate_content(text, num_slides=5, language=Language.CHINESE)
        last = slides[-1]
        assert last.title == "\u6838\u5fc3\u8981\u70b9"

    def test_mock_no_summary_for_few_slides(self):
        """2 slides should not have a separate summary."""
        text = "Short content for testing."
        slides = mock_generate_content(text, num_slides=2)
        assert len(slides) == 2

    def test_mock_summary_has_content(self):
        """Summary slide should have content points from other slides."""
        text = "First topic is important. Second topic matters too. Third point is key. " * 10
        slides = mock_generate_content(text, num_slides=5)
        last = slides[-1]
        assert len(last.content) >= 1

    def test_mock_improved_sentence_splitting(self):
        """Verify new separators like semicolons and ellipses are handled."""
        text = "Point one; Point two\u2026 Point three\uff01"
        slides = mock_generate_content(text, num_slides=2)
        assert len(slides) == 2


class TestPresentationConfigOverlay:
    """Tests for PresentationConfig overlay_opacity field."""

    def test_config_default_overlay(self):
        config = PresentationConfig()
        assert config.overlay_opacity == 130

    def test_config_custom_overlay(self):
        config = PresentationConfig(overlay_opacity=200)
        assert config.overlay_opacity == 200


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
