"""Tests for the document parser module."""
import io
import pytest
from docx import Document

from src.parser import (
    parse_document,
    extract_text_from_docx,
    extract_text_from_txt,
    SUPPORTED_EXTENSIONS,
)


class TestExtractTextFromDocx:
    def _make_docx(self, paragraphs, headings=None):
        """Helper to create an in-memory DOCX document."""
        doc = Document()
        if headings:
            for level, text in headings:
                doc.add_heading(text, level)
        for text in paragraphs:
            doc.add_paragraph(text)
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf

    def test_basic_extraction(self):
        buf = self._make_docx(["Hello World from Docx"])
        text = extract_text_from_docx(buf)
        assert "Hello World from Docx" in text

    def test_multiple_paragraphs(self):
        paragraphs = ["First paragraph", "Second paragraph", "Third paragraph"]
        buf = self._make_docx(paragraphs)
        text = extract_text_from_docx(buf)
        for p in paragraphs:
            assert p in text

    def test_empty_paragraphs_skipped(self):
        doc = Document()
        doc.add_paragraph("Content")
        doc.add_paragraph("")  # Empty
        doc.add_paragraph("   ")  # Whitespace only
        doc.add_paragraph("More Content")
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        text = extract_text_from_docx(buf)
        assert "Content" in text
        assert "More Content" in text

    def test_headings_preserved(self):
        buf = self._make_docx([], headings=[(1, "Main Title")])
        text = extract_text_from_docx(buf)
        assert "# Main Title" in text


class TestExtractTextFromTxt:
    def test_utf8_text(self):
        content = "Hello, World!\nThis is a test."
        buf = io.BytesIO(content.encode("utf-8"))
        text = extract_text_from_txt(buf)
        assert text == content

    def test_latin1_fallback(self):
        content = "caf\u00e9"
        buf = io.BytesIO(content.encode("latin-1"))
        text = extract_text_from_txt(buf)
        assert "caf" in text


class TestParseDocument:
    def test_docx_parsing(self):
        doc = Document()
        doc.add_paragraph("Test content for parsing")
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        buf.name = "test.docx"
        text = parse_document(buf)
        assert "Test content for parsing" in text

    def test_txt_parsing(self):
        content = "Plain text content for testing"
        buf = io.BytesIO(content.encode("utf-8"))
        buf.name = "test.txt"
        text = parse_document(buf)
        assert text == content

    def test_markdown_parsing(self):
        content = "# Heading\n\nSome markdown content"
        buf = io.BytesIO(content.encode("utf-8"))
        buf.name = "test.md"
        text = parse_document(buf)
        assert "Heading" in text

    def test_unsupported_format_raises(self):
        buf = io.BytesIO(b"data")
        buf.name = "test.xyz"
        with pytest.raises(ValueError, match="Unsupported file type"):
            parse_document(buf)

    def test_no_name_attribute_raises(self):
        buf = io.BytesIO(b"data")
        with pytest.raises(ValueError, match="no name attribute"):
            parse_document(buf)

    def test_supported_extensions_complete(self):
        assert "docx" in SUPPORTED_EXTENSIONS
        assert "pdf" in SUPPORTED_EXTENSIONS
        assert "txt" in SUPPORTED_EXTENSIONS
        assert "md" in SUPPORTED_EXTENSIONS


if __name__ == "__main__":
    pytest.main([__file__, "-v"])
